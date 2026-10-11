"""Minimal Markdown to .docx converter, because pandoc and LibreOffice are both
absent on this machine.

Handles what these deliverables actually use: ATX headings, paragraphs, bullet
and numbered lists, pipe tables, fenced code blocks, blockquotes, horizontal
rules, and inline **bold** / *italic* / `code`. Anything else passes through as
plain text rather than silently disappearing.

Use:
    from md2docx import convert
    convert("audit_report.md", "audit_report.docx")
or
    python md2docx.py in.md out.docx
    python md2docx.py --demo        # self-check
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)")


def _add_runs(par, text):
    """Split text on inline markup and add styled runs."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            r.font.name = "Menlo"
            r.font.size = Pt(9)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def _is_table_sep(line):
    s = line.strip()
    return bool(s) and set(s) <= set("|-: ") and "-" in s and "|" in s


def _split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def convert(md_path, docx_path, title=None):
    md_path = os.path.abspath(md_path)
    with open(md_path, encoding="utf-8") as fh:
        lines = fh.read().split("\n")

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1.0)
        section.top_margin = section.bottom_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(6)

    if title:
        doc.core_properties.title = title

    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        # fenced code
        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            r = p.add_run("\n".join(buf))
            r.font.name = "Menlo"
            r.font.size = Pt(8.5)
            continue

        # table
        if s.startswith("|") and i + 1 < len(lines) and _is_table_sep(lines[i + 1]):
            header = _split_row(line)
            i += 2
            body = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                body.append(_split_row(lines[i]))
                i += 1
            ncol = len(header)
            t = doc.add_table(rows=1, cols=ncol)
            t.style = "Light Grid Accent 1"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, txt in zip(t.rows[0].cells, header):
                c.paragraphs[0].clear()
                _add_runs(c.paragraphs[0], txt)
                for r in c.paragraphs[0].runs:
                    r.bold = True
                    r.font.size = Pt(9)
            for row in body:
                cells = t.add_row().cells
                for c, txt in zip(cells, row[:ncol]):
                    c.paragraphs[0].clear()
                    _add_runs(c.paragraphs[0], txt)
                    for r in c.paragraphs[0].runs:
                        r.font.size = Pt(9)
            doc.add_paragraph()
            continue

        # image:  ![alt](relative/path.png)
        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)\s*$", s)
        if m:
            rel = m.group(1)
            img = os.path.normpath(os.path.join(os.path.dirname(md_path), rel))
            if os.path.exists(img):
                from PIL import Image as _I
                w_px, h_px = _I.open(img).size
                # full text width, capped so tall figures still fit a page
                width_in = 6.5
                if h_px / w_px * width_in > 8.0:
                    width_in = 8.0 * w_px / h_px
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(img, width=Inches(width_in))
            else:
                p_ = doc.add_paragraph()
                p_.add_run(f"[missing figure: {rel}]").italic = True
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            h = doc.add_heading(level=min(len(m.group(1)), 4))
            _add_runs(h, m.group(2))
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
            i += 1
            continue

        # horizontal rule
        if s in ("---", "***", "___"):
            p = doc.add_paragraph()
            p.add_run("_" * 78).font.color.rgb = RGBColor(0xA0, 0xA0, 0xA0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # blockquote
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            _add_runs(p, s.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # bullet
        m = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if m:
            indent = (len(line) - len(line.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.25 + 0.25 * indent)
            _add_runs(p, m.group(1))
            i += 1
            continue

        # numbered. Word's "List Number" style renumbers independently of the
        # source and continues across every numbered list in the document, so a
        # reference list rendered that way stops matching its in-text citations.
        # Keep the literal number from the markdown instead.
        m = re.match(r"^\s*(\d+)[.)]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.33)
            p.paragraph_format.first_line_indent = Inches(-0.33)
            p.paragraph_format.space_after = Pt(3)
            _add_runs(p, f"{m.group(1)}. {m.group(2)}")
            i += 1
            continue

        if not s:
            i += 1
            continue

        # paragraph: join continuation lines
        buf = [s]
        j = i + 1
        while j < len(lines):
            nxt = lines[j]
            ns = nxt.strip()
            if (not ns or ns.startswith(("#", "|", ">", "```", "---"))
                    or re.match(r"^\s*([-*+]|\d+[.)])\s+", nxt)):
                break
            buf.append(ns)
            j += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _add_runs(p, " ".join(buf))
        i = j

    doc.save(docx_path)
    return docx_path


def _demo():
    """Self-check: every construct survives the round trip into a saved file."""
    import os
    import tempfile
    src = """# H1
Intro **bold** and *italic* and `code`.

## H2
- bullet one
- bullet two

1. first
2. second

| a | b |
| --- | --- |
| 1 | 2 |

> a quote

```
literal
```
---
tail paragraph
"""
    d = tempfile.mkdtemp()
    md, dx = os.path.join(d, "t.md"), os.path.join(d, "t.docx")
    with open(md, "w", encoding="utf-8") as fh:
        fh.write(src)
    convert(md, dx)
    doc = Document(dx)
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert heads == ["H1", "H2"], heads
    assert len(doc.tables) == 1 and doc.tables[0].rows[0].cells[0].text == "a"
    assert any(r.bold for p in doc.paragraphs for r in p.runs), "no bold run"
    assert any(p.style.name == "List Bullet" for p in doc.paragraphs)
    assert any(re.match(r"\s*\d+\.", p.text) for p in doc.paragraphs)
    assert any("literal" in p.text for p in doc.paragraphs)
    assert any("tail paragraph" in p.text for p in doc.paragraphs)
    assert os.path.getsize(dx) > 8000
    print("md2docx self-check passed")


if __name__ == "__main__":
    if len(sys.argv) == 2 and sys.argv[1] == "--demo":
        _demo()
    elif len(sys.argv) >= 3:
        print(convert(sys.argv[1], sys.argv[2]))
    else:
        print(__doc__)
        sys.exit(1)
